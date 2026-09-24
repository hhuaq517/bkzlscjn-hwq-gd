# -*- coding: utf-8 -*-
"""docx → PDF 离线生成器（reportlab，两遍编译回填目录页码）。"""
import io
import os
import re
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Frame, Image,
                                KeepTogether, PageBreak, PageTemplate,
                                Paragraph as RLPara, Spacer, Table as RLTable,
                                TableStyle)

SRC = sys.argv[1] if len(sys.argv) > 1 else "gk_new.docx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "out.pdf"
TMPDIR = os.path.dirname(os.path.abspath(SRC))

PW, PH = A4
M_T, M_B, M_L, M_R = 2.2 * cm, 2.2 * cm, 2.4 * cm, 2.4 * cm
CW = PW - M_L - M_R                       # 459.2pt
HDR = "2027 年国考考情手册"
SIGN = "广东中公教研"
NUMFMT = {}                               # numId → 'decimal' | 'bullet'

INK = colors.HexColor("#1A1A1A")
RED = colors.HexColor("#A62B2B")
GRAY = colors.HexColor("#6E6E6E")
LGRAY = colors.HexColor("#9E9E9E")
BORDER = colors.HexColor("#D6D6D6")

pdfmetrics.registerFont(TTFont("Song", "/System/Library/Fonts/Supplemental/Songti.ttc", subfontIndex=0))
pdfmetrics.registerFont(TTFont("Hei", "/System/Library/Fonts/STHeiti Medium.ttc", subfontIndex=0))
pdfmetrics.registerFont(TTFont("HeiL", "/System/Library/Fonts/STHeiti Light.ttc", subfontIndex=0))
pdfmetrics.registerFontFamily("Song", normal="Song", bold="Hei", italic="Song", boldItalic="Hei")


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def rgb_of(rPr):
    if rPr is None:
        return None
    c = rPr.find(qn("w:color"))
    if c is None:
        return None
    v = c.get(qn("w:val"))
    return ("#" + v.upper()) if v and v != "auto" else None


def run_markup(r, def_size, cut_tab=False):
    t = r.text
    if not t:
        return ""
    if cut_tab and "\t" in t:
        t = t.split("\t")[0]
        if not t:
            return ""
    t = t.replace("\u2212", "-")           # U+2212/U+2013 在 Songti 子集中均无字形
    rPr = r._r.find(qn("w:rPr"))
    bold = rPr is not None and rPr.find(qn("w:b")) is not None
    sz = None
    if rPr is not None:
        s = rPr.find(qn("w:sz"))
        if s is not None:
            sz = int(s.get(qn("w:val"))) / 2.0
    col = rgb_of(rPr)
    body = esc(t).replace("\t", " ")
    if body.startswith(" "):
        body = "&nbsp;" + body[1:]
    if body.endswith(" "):
        body = body[:-1] + "&nbsp;"
    attrs = ""
    if col:
        attrs += f' color="{col}"'
    if sz and abs(sz - def_size) > 0.25:
        attrs += f' size="{sz}"'
    if attrs:
        body = f"<font{attrs}>{body}</font>"
    if bold:
        body = f"<b>{body}</b>"
    return body


def para_markup(p, def_size, cut_tab=False):
    s = "".join(run_markup(r, def_size, cut_tab) for r in p.runs)
    return re.sub(r"[\s\u3000]+$", "", s)


def max_size(p, fallback):
    """段落内最大字号"""
    out = fallback
    for r in p.runs:
        rPr = r._r.find(qn("w:rPr"))
        if rPr is not None:
            s = rPr.find(qn("w:sz"))
            if s is not None:
                v = int(s.get(qn("w:val"))) / 2.0
                if out is fallback or v > out:
                    out = v
    return out


def p_get(p, tag):
    pPr = p._p.find(qn("w:pPr"))
    return pPr.find(qn(tag)) if pPr is not None else None


def p_jc(p):
    j = p_get(p, "w:jc")
    return j.get(qn("w:val")) if j is not None else None


def p_num(p):
    n = p_get(p, "w:numPr")
    if n is None:
        return None
    nid = n.find(qn("w:numId"))
    return int(nid.get(qn("w:val"))) if nid is not None else None


def load_numbering(src):
    z = zipfile.ZipFile(src)
    if "word/numbering.xml" not in z.namelist():
        return
    nb = z.read("word/numbering.xml").decode()
    n2a = dict(re.findall(r'<w:num w:numId="(\d+)"[^>]*>\s*<w:abstractNumId w:val="(\d+)"/>', nb))
    amap = {}
    for m in re.finditer(r'<w:abstractNum w:abstractNumId="(\d+)"[^>]*>(.*?)</w:abstractNum>', nb, re.S):
        lvl0 = re.search(r'<w:lvl w:ilvl="0".*?<w:numFmt w:val="(\w+)"/>', m.group(2), re.S)
        if lvl0:
            amap[m.group(1)] = lvl0.group(1)
    for nid, aid in n2a.items():
        NUMFMT[int(nid)] = amap.get(aid, "bullet")


def load_images(src):
    """rId → (bytes, 显示宽pt, 显示高pt)"""
    z = zipfile.ZipFile(src)
    d = z.read("word/document.xml").decode()
    rels = z.read("word/_rels/document.xml.rels").decode()
    rmap = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    ext_map = {}
    for m in re.finditer(r'<wp:extent cx="(\d+)" cy="(\d+)"/>.*?r:embed="([^"]+)"', d, re.S):
        w, h, rid = int(m.group(1)) / 360000 * 28.3465, int(m.group(2)) / 360000 * 28.3465, m.group(3)
        ext_map[rid] = (w, h)
    out = {}
    for rid, tgt in rmap.items():
        if "media/" in tgt and tgt.endswith(".png"):
            full = "word/" + tgt
            out[rid] = (z.read(full), *ext_map.get(rid, (0, 0)))
    return out


def draw_images(src, outdir):
    z = zipfile.ZipFile(src)
    for n in z.namelist():
        if n.startswith("word/media/") and n.endswith(".png"):
            with open(os.path.join(outdir, os.path.basename(n)), "wb") as f:
                f.write(z.read(n))


class HB(BaseDocTemplate):
    def __init__(self, fn, **kw):
        super().__init__(fn, **kw)
        self.h1_pages = {}

    def afterFlowable(self, fl):
        t = getattr(fl, "_h1_title", None)
        if t and t not in self.h1_pages:
            self.h1_pages[t] = self.page


def on_page(canv, doc):
    n = canv.getPageNumber()
    if n == 1:                                # 封面无页眉页脚
        return
    canv.saveState()
    y = PH - M_T + 0.62 * cm
    canv.setFont("HeiL", 8)
    canv.setFillColor(GRAY)
    canv.drawString(M_L, y, HDR)
    canv.drawRightString(PW - M_R, y, SIGN)
    canv.setStrokeColor(colors.HexColor("#D6D6D6"))
    canv.setLineWidth(0.5)
    canv.line(M_L, y - 5.5, PW - M_R, y - 5.5)
    canv.setFont("Song", 8.4)
    canv.setFillColor(LGRAY)
    canv.drawCentredString(PW / 2, M_B - 0.52 * cm, f"— {n} —")
    canv.restoreState()


def para_style(p, size, line, S, align=None, leading_k=1.16):
    """按段落内最大字号生成样式"""
    ms = max_size(p, None)
    fs = ms if ms else size
    st = ParagraphStyle("x", fontName="Song", fontSize=fs,
                        leading=fs * leading_k * (line if line < 1.9 else 1.0),
                        textColor=INK, alignment=align or TA_LEFT, spaceBefore=0, spaceAfter=0)
    return st


def cell_paras(cell, size, line, S):
    """单元格 → (markup, ParagraphStyle) 列表"""
    out = []
    for p in cell.paragraphs:
        mk = para_markup(p, size)
        if not mk:
            continue
        jc = p_jc(p)
        align = TA_CENTER if jc == "center" else (TA_RIGHT if jc == "right" else TA_LEFT)
        st = ParagraphStyle("c", fontName="Song", fontSize=size, leading=size * 1.32,
                            textColor=INK, alignment=align)
        out.append((mk, st))
    return out


def build_table(t, S, size, line):
    ncol = len(t.columns)
    grid = [int(g.get(qn("w:w"))) for g in t._tbl.iter(qn("w:gridCol"))]
    scale = CW / sum(grid) if grid and sum(grid) > 0 else CW / ncol
    widths = [g * scale for g in grid] if grid else [CW / ncol] * ncol
    # 单元格字号：取整表 run 众数
    szs = []
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    rPr = run._r.find(qn("w:rPr"))
                    if rPr is not None:
                        s = rPr.find(qn("w:sz"))
                        if s is not None:
                            szs.append(int(s.get(qn("w:val"))) / 2.0)
    csz = max(set(szs), key=szs.count) if szs else 8.8
    data, style = [], []
    seen = set()
    for ri, row in enumerate(t.rows):
        rdata, ci = [], 0
        for cell in row.cells:
            if cell._tc in seen:                # 合并单元格重复引用
                ci += 1
                continue
            seen.add(cell._tc)
            paras = cell_paras(cell, csz, line, S)
            if not paras:
                rdata.append("")
                ci += 1
                continue
            fl = [RLPara(mk, st) for mk, st in paras]
            rdata.append(fl if len(fl) > 1 else fl[0])
            ci += 1
        while len(rdata) < ncol:
            rdata.append("")
        data.append(rdata)
    tbl = RLTable(data, colWidths=widths, repeatRows=1)
    cmds = [
        ("GRID", (0, 0), (-1, -1), 0.5, BORDER),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]
    if len(t.rows) > 1:                       # 表头行：浅灰底 + 加粗
        cmds += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F2F2"))]
        for ci in range(ncol):
            v = data[0][ci]
            if isinstance(v, RLPara):
                v.style.fontName = "Hei"
    tbl.setStyle(TableStyle(cmds))
    return tbl


def build(src, out, toc_pages=None):
    doc_xml = Document(src)
    size, line = styles_size_line(doc_xml)
    S = make_styles(size, line)
    load_numbering(src)
    draw_images(src, TMPDIR)
    IMG = load_images(src)

    num_ctr = {}
    story = []
    h1_seen = set()
    blocks = []
    for child in doc_xml.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(("p", Paragraph(child, doc_xml._body)))
        elif child.tag == qn("w:tbl"):
            blocks.append(("tbl", Table(child, doc_xml._body)))

    # 切分：封面 = 前到第一个分页符；目录 = 之后到首个 H1
    cover_end = next(i for i, (k, o) in enumerate(blocks)
                     if k == "p" and any(b.get(qn("w:type")) == "page"
                                         for b in o._p.iter(qn("w:br"))))
    toc_start = cover_end + 1
    first_h1 = next(i for i, (k, o) in enumerate(blocks)
                    if k == "p" and o.style.name == "Heading 1")
    toc_end = first_h1

    cover = blocks[:cover_end]
    toc = blocks[toc_start:toc_end]
    body = blocks[toc_end:]

    # ---- 封面 ----
    for k, o in cover:
        if k != "p":
            continue
        mk = para_markup(o, size)
        if not mk:
            continue
        st = para_style(o, size, line, S, TA_CENTER, leading_k=1.25)
        story.append(Spacer(1, 6))
        story.append(RLPara(mk, st))
    # 封面数据卡表
    for i, (k, o) in enumerate(cover):
        if k == "tbl":
            t = build_table(o, S, size, line)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F5F5F5")),
                                   ("BACKGROUND", (0, 0), (-1, -1), colors.transparent)]))
            story.append(Spacer(1, 10))
            story.append(t)
    # 封底署名区
    sig = [o for k, o in cover if k == "p" and o.text.strip()]
    story.append(Spacer(1, 0))

    # ---- 目录 ----
    story.append(PageBreak())
    for k, o in toc:
        if k != "p":
            continue
        t = o.text.strip()
        if not t:
            continue
        if t.startswith("目"):
            mk = para_markup(o, size)
            st = ParagraphStyle("tocT", fontName="Hei", fontSize=15, leading=22,
                                textColor=INK, spaceAfter=4)
            story.append(RLPara(mk, st))
            story.append(Spacer(1, 1))
            story.append(HRule(CW, RED, 1.1))
            story.append(Spacer(1, 10))
            continue
        parts = t.split("\t")
        left = para_markup(o, 10.2, cut_tab=True)
        left = re.sub(r"\s+$", "", left)
        page = toc_pages.get(parts[0].strip(), "—") if toc_pages else (parts[1] if len(parts) > 1 else "—")
        print(f"TOCDBG {parts[0].strip()[:12]!r} tab={len(parts)} page={page!r}")
        row = RLTable([[RLPara(left, S["toc"]), RLPara(str(page), S["tocp"])]],
                      colWidths=[CW - 34, 34])
        row.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 1.2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2),
            ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#EAEAEA")),
        ]))
        story.append(row)

    # ---- 正文 ----
    for k, o in body:
        if k == "p":
            p = o
            has_img = p._p.findall(".//" + qn("w:drawing"))
            t = p.text.strip()
            if has_img:
                story.append(Spacer(1, 6))
                grp = []
                for rid, (blob, w, h) in IMG.items():
                    pass                       # 位置无关，仅按文档顺序
                for run in p.runs:
                    for blip in run._r.iter(qn("a:blip")):
                        rid = blip.get(qn("r:embed"))
                        if rid in IMG:
                            blob, w, h = IMG[rid]
                            iw = min(w, CW)
                            ih = h * (iw / w) if w else 0
                            fn = os.path.join(TMPDIR, "img_%d.png" % len(grp))
                            with open(fn, "wb") as f:
                                f.write(blob)
                            grp.append(Image(fn, width=iw, height=ih))
                if grp:
                    inner = KeepTogether(grp + [Spacer(1, 2)])
                    story.append(inner)
                continue
            if not t:
                continue
            stname = p.style.name
            mk = para_markup(o, size)
            if not mk:
                continue
            if stname == "Heading 1":
                if t in h1_seen:
                    continue
                h1_seen.add(t)
                fl = RLPara(mk, S["h1"])
                fl._h1_title = t
                story.append(CondPageBreak(72))
                story.append(fl)
                story.append(HRule(CW, RED, 0.8))
                story.append(Spacer(1, 7))
            elif stname == "Heading 2":
                story.append(CondPageBreak(52))
                story.append(RLPara(mk, S["h2"]))
            elif re.match(r"^图\s*[\dA-Z]+-\d+", t):
                story.append(RLPara(mk, S["cap"]))
            else:
                num = p_num(p)
                jc = p_jc(p)
                if num is not None:
                    fmt = NUMFMT.get(num, "bullet")
                    if fmt == "decimal":
                        num_ctr[num] = num_ctr.get(num, 0) + 1
                        mk = f"<b><font color=\"#A62B2B\">{num_ctr[num]}.</font></b> " + mk
                    else:
                        mk = f"<font color=\"#A62B2B\">·</font> " + mk
                    story.append(RLPara(mk, S["li"]))
                else:
                    num_ctr.clear()
                    al = TA_CENTER if jc == "center" else TA_LEFT
                    st = S["body"] if al == TA_LEFT else ParagraphStyle(
                        "bc", parent=S["body"], alignment=al)
                    story.append(RLPara(mk, st))
        else:
            story.append(Spacer(1, 4))
            story.append(build_table(o, S, size, line))
            story.append(Spacer(1, 8))

    doc = HB(out, pagesize=A4, leftMargin=M_L, rightMargin=M_R,
             topMargin=M_T, bottomMargin=M_B,
             title="2027 年国考考情手册", author=SIGN)
    fr = Frame(M_L, M_B, CW, PH - M_T - M_B, id="f",
               leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="pt", frames=[fr], onPage=on_page)])
    doc.build(story)
    return doc.h1_pages


def styles_size_line(doc):
    dd = doc.styles.element.find(qn("w:docDefaults"))
    size, line = 10.5, 1.5
    if dd is not None:
        rPr = dd.find(qn("w:rPrDefault"))
        if rPr is not None:
            rp = rPr.find(qn("w:rPr"))
            if rp is not None:
                s = rp.find(qn("w:sz"))
                if s is not None:
                    size = int(s.get(qn("w:val"))) / 2.0
        pPr = dd.find(qn("w:pPrDefault"))
        if pPr is not None:
            pp = pPr.find(qn("w:pPr"))
            if pp is not None:
                sp = pp.find(qn("w:spacing"))
                if sp is not None and sp.get(qn("w:lineRule")) == "auto" and sp.get(qn("w:line")):
                    line = int(sp.get(qn("w:line"))) / 240.0
    return size, line


def make_styles(size, line):
    lead = size * 1.16 * line
    base = dict(fontName="Song", fontSize=size, leading=lead, textColor=INK,
                spaceBefore=0, spaceAfter=size * 0.52)
    S = {}
    S["body"] = ParagraphStyle("body", **base)
    S["h1"] = ParagraphStyle("h1", fontName="Hei", fontSize=15, leading=21,
                             textColor=RED, spaceBefore=4, spaceAfter=6)
    S["h2"] = ParagraphStyle("h2", fontName="Hei", fontSize=11.4, leading=16.5,
                             textColor=INK, spaceBefore=7, spaceAfter=4)
    S["cap"] = ParagraphStyle("cap", fontName="HeiL", fontSize=8.6, leading=12.4,
                              textColor=GRAY, alignment=TA_CENTER,
                              spaceBefore=3, spaceAfter=9)
    S["li"] = ParagraphStyle("li", fontName="Song", fontSize=size, leading=lead,
                             textColor=INK, leftIndent=15, firstLineIndent=-11,
                             spaceAfter=size * 0.4)
    S["toc"] = ParagraphStyle("toc", fontName="Hei", fontSize=10.2, leading=15,
                              textColor=INK)
    S["tocp"] = ParagraphStyle("tocp", fontName="Song", fontSize=9.6, leading=15,
                               textColor=GRAY, alignment=TA_RIGHT)
    return S


from reportlab.platypus import Flowable


class HRule(Flowable):
    def __init__(self, w, col, th):
        Flowable.__init__(self)
        self.w, self.col, self.th = w, col, th
        self.height = th + 3

    def draw(self):
        self.canv.setStrokeColor(self.col)
        self.canv.setLineWidth(self.th)
        self.canv.line(0, self.th + 1.5, self.w, self.th + 1.5)


def main():
    global TMPDIR
    pages = None
    for it in range(1, 5):
        pages = build(SRC, OUT if it > 1 else "pass1.pdf", pages)
        print(f"第 {it} 遍：{len(pages)} 章，附录C 落 p{list(pages.values())[-1]}")
        if it > 1 and pages == getattr(main, "_prev", None):
            print("  → 页码已收敛")
            break
        main._prev = pages
    print("\n已保存:", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
